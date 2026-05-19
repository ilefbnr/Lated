from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# ---- Theme ----
NAVY    = RGBColor(0x0B, 0x1F, 0x3A)   # background
ACCENT  = RGBColor(0x00, 0xB8, 0xD9)   # cyan accent
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
MUTED   = RGBColor(0xB8, 0xC4, 0xD6)
GREEN   = RGBColor(0x36, 0xC2, 0x7A)
AMBER   = RGBColor(0xF5, 0xA6, 0x23)
GREY    = RGBColor(0x7A, 0x86, 0x99)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

blank = prs.slide_layouts[6]

# ---------- helpers ----------
def add_bg(slide, color=NAVY):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_accent_bar(slide, x, y, w, h, color=ACCENT):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    bar.fill.solid(); bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar

def add_text(slide, x, y, w, h, text, *, size=18, bold=False,
             color=WHITE, align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb

def add_bullets(slide, x, y, w, h, items, *, size=16, color=WHITE,
                bullet_color=ACCENT, line_spacing=1.25):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        r1 = p.add_run(); r1.text = "▸  "
        r1.font.size = Pt(size); r1.font.bold = True
        r1.font.color.rgb = bullet_color
        r2 = p.add_run(); r2.text = item
        r2.font.size = Pt(size); r2.font.color.rgb = color
        r2.font.name = "Calibri"
    return tb

def add_chip(slide, x, y, w, h, text, *, fill=ACCENT, color=NAVY,
             size=11, bold=True):
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    chip.fill.solid(); chip.fill.fore_color.rgb = fill
    chip.line.fill.background()
    tf = chip.text_frame
    tf.margin_left = Emu(60000); tf.margin_right = Emu(60000)
    tf.margin_top = Emu(20000); tf.margin_bottom = Emu(20000)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.name = "Calibri"
    return chip

def add_footer(slide, text):
    add_text(slide, Inches(0.5), Inches(7.05),
             Inches(12.3), Inches(0.35),
             text, size=10, color=MUTED)

# =====================================================================
# SLIDE 1 — Identity
# =====================================================================
s1 = prs.slides.add_slide(blank)
add_bg(s1)

# left accent column
add_accent_bar(s1, 0, 0, Inches(0.35), SH, color=ACCENT)

# Photo placeholder (left)
photo_x, photo_y = Inches(1.1), Inches(1.6)
photo_w, photo_h = Inches(4.2), Inches(4.2)
photo = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, photo_x, photo_y, photo_w, photo_h)
photo.fill.solid(); photo.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
photo.line.color.rgb = ACCENT
photo.line.width = Pt(1.5)
ptf = photo.text_frame
ptf.margin_left = Emu(0); ptf.margin_right = Emu(0)
p = ptf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "\n\n[ INSERT PHOTO ]"
r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = MUTED

p2 = ptf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run(); r2.text = "(replace this placeholder)"
r2.font.size = Pt(11); r2.font.color.rgb = GREY; r2.font.italic = True

# Right side — name and university
right_x = Inches(6.2)
add_chip(s1, right_x, Inches(1.6), Inches(2.1), Inches(0.4),
         "PFE 2026", fill=ACCENT, color=NAVY, size=11)

add_text(s1, right_x, Inches(2.15), Inches(6.5), Inches(1.1),
         "Ilef Bennour", size=44, bold=True, color=WHITE)

# divider
div = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                         right_x, Inches(3.25), Inches(0.7), Emu(38100))
div.fill.solid(); div.fill.fore_color.rgb = ACCENT
div.line.fill.background()

add_text(s1, right_x, Inches(3.45), Inches(6.5), Inches(0.45),
         "Intern  •  End-of-Studies Project", size=15, color=MUTED)

add_text(s1, right_x, Inches(4.15), Inches(6.5), Inches(0.55),
         "ENSI", size=24, bold=True, color=WHITE)
add_text(s1, right_x, Inches(4.70), Inches(6.5), Inches(0.45),
         "École Nationale des Sciences de l’Informatique", size=14, color=MUTED)
add_text(s1, right_x, Inches(5.10), Inches(6.5), Inches(0.45),
         "Université de la Manouba", size=14, color=MUTED)

# Bottom project name
add_text(s1, Inches(1.1), Inches(6.30), Inches(11.0), Inches(0.5),
         "LateD — Lateral Movement Detection Platform",
         size=16, bold=True, color=ACCENT)

add_footer(s1, "PFE 2026 • ENSI — Université de la Manouba")

# =====================================================================
# SLIDE 2 — Project description
# =====================================================================
s2 = prs.slides.add_slide(blank)
add_bg(s2)
add_accent_bar(s2, 0, 0, SW, Inches(0.12), color=ACCENT)

# title
add_text(s2, Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.6),
         "Project Description", size=30, bold=True, color=WHITE)
add_text(s2, Inches(0.6), Inches(0.95), Inches(12.0), Inches(0.4),
         "LateD — SOC platform for real-time lateral movement detection",
         size=15, color=ACCENT)

# === LEFT: General idea ===
left_x = Inches(0.6); top = Inches(1.65)
card1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            left_x, top, Inches(6.0), Inches(5.1))
card1.fill.solid(); card1.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
card1.line.color.rgb = ACCENT; card1.line.width = Pt(0.75)

add_chip(s2, left_x + Inches(0.3), top + Inches(0.3),
         Inches(1.5), Inches(0.4), "GENERAL IDEA",
         fill=ACCENT, color=NAVY, size=11)

add_text(s2, left_x + Inches(0.3), top + Inches(0.85),
         Inches(5.5), Inches(0.9),
         "Detect lateral movement inside enterprise networks "
         "using Temporal Graph Neural Networks.",
         size=15, bold=True, color=WHITE)

add_bullets(s2, left_x + Inches(0.3), top + Inches(2.0),
            Inches(5.5), Inches(3.0),
            [
                "Lateral Movement detection via TGNN",
                "Internal reconnaissance (heuristic + behavioral)",
                "Attack-path reconstruction & correlation",
                "Real-time SOC dashboard — command-center UX",
                "Explainable scoring — every alert is traceable",
            ],
            size=13, color=WHITE, bullet_color=ACCENT)

# === RIGHT: Execution plan ===
right_x = Inches(6.85)
card2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            right_x, top, Inches(6.0), Inches(5.1))
card2.fill.solid(); card2.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
card2.line.color.rgb = ACCENT; card2.line.width = Pt(0.75)

add_chip(s2, right_x + Inches(0.3), top + Inches(0.3),
         Inches(1.9), Inches(0.4), "EXECUTION PLAN",
         fill=ACCENT, color=NAVY, size=11)

phases = [
    ("1", "OFFLINE phase", "Train TGNN on LANL flows + redteam ground truth.", ACCENT),
    ("2", "ONLINE phase",  "Real-time detection on PCAP / Zeek / NetFlow streams.", ACCENT),
    ("3", "Correlation",   "Chain recon → LM, reconstruct attack paths.", AMBER),
    ("4", "Supervision",   "FastAPI + WebSocket + Next.js SOC dashboard.", GREEN),
]

py = top + Inches(0.9)
for num, name, desc, col in phases:
    # number circle
    circ = s2.shapes.add_shape(MSO_SHAPE.OVAL,
                               right_x + Inches(0.3), py,
                               Inches(0.55), Inches(0.55))
    circ.fill.solid(); circ.fill.fore_color.rgb = col
    circ.line.fill.background()
    ctf = circ.text_frame
    ctf.margin_left = Emu(0); ctf.margin_right = Emu(0)
    ctf.margin_top = Emu(0); ctf.margin_bottom = Emu(0)
    cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.CENTER
    cr = cp.add_run(); cr.text = num
    cr.font.size = Pt(16); cr.font.bold = True; cr.font.color.rgb = NAVY

    add_text(s2, right_x + Inches(1.05), py - Inches(0.02),
             Inches(4.8), Inches(0.35),
             name, size=14, bold=True, color=WHITE)
    add_text(s2, right_x + Inches(1.05), py + Inches(0.30),
             Inches(4.8), Inches(0.5),
             desc, size=11, color=MUTED)
    py += Inches(0.95)

add_footer(s2, "Slide 2 / 4  •  LateD — Project description")

# =====================================================================
# SLIDE 3 — Training methods comparison
# =====================================================================
s_train = prs.slides.add_slide(blank)
add_bg(s_train)
add_accent_bar(s_train, 0, 0, SW, Inches(0.12), color=ACCENT)

add_text(s_train, Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.6),
         "LM Detector — Training methods", size=30, bold=True, color=WHITE)
add_text(s_train, Inches(0.6), Inches(0.95), Inches(12.0), Inches(0.4),
         "Three regimes, same TGNN architecture, comparable on eval split.",
         size=15, color=ACCENT)

# ---- Three regime cards ----
card_top  = Inches(1.65)
card_h    = Inches(3.7)
card_w    = Inches(4.05)
card_xs   = [Inches(0.6), Inches(4.85), Inches(9.10)]
card_clrs = [ACCENT, AMBER, GREEN]
card_tags = ["A · SELF-SUPERVISED", "B · SEMI-SUPERVISED",
             "C · FULLY SUPERVISED"]
card_titles = [
    "SSL link prediction",
    "SSL backbone + supervised head",
    "End-to-end supervised",
]
card_files = [
    "pretrain_ssl.py",
    "pretrain_ssl.py  →  finetune_lm.py",
    "supervised_lm.py",
]
card_bullets = [
    [
        "Trains TGN backbone only — no labels used",
        "Loss: link-prediction on benign edges",
        "Score = 1 − P(link) at inference",
        "Pros: needs no red-team labels",
        "Cons: anomaly ≠ lateral movement",
    ],
    [
        "Backbone frozen from regime A",
        "LM head trained on weak labels (weak_labeler.py)",
        "pos_weight handles ≈1:333 imbalance",
        "Pros: little label cost, strong reuse",
        "Cons: limited by SSL representation",
    ],
    [
        "Backbone + head trained jointly",
        "Direct BCE on LM labels (mask recon)",
        "Backbone gradient driven by LM signal",
        "Pros: best-case ceiling for this data",
        "Cons: needs labels for the whole stream",
    ],
]

for x, color, tag, title, fname, bullets in zip(
        card_xs, card_clrs, card_tags, card_titles, card_files, card_bullets):
    card = s_train.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, x, card_top, card_w, card_h,
    )
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
    card.line.color.rgb = color; card.line.width = Pt(1)

    add_chip(s_train, x + Inches(0.25), card_top + Inches(0.22),
             Inches(2.4), Inches(0.4), tag, fill=color, color=NAVY, size=10)
    add_text(s_train, x + Inches(0.25), card_top + Inches(0.75),
             card_w - Inches(0.5), Inches(0.45),
             title, size=15, bold=True, color=WHITE)
    add_text(s_train, x + Inches(0.25), card_top + Inches(1.18),
             card_w - Inches(0.5), Inches(0.35),
             fname, size=10, color=MUTED, font="Consolas")
    add_bullets(s_train, x + Inches(0.25), card_top + Inches(1.60),
                card_w - Inches(0.5), card_h - Inches(1.75),
                bullets, size=11, color=WHITE,
                bullet_color=color, line_spacing=1.25)

# ---- Comparison table (eval split, illustrative slots) ----
tbl_top = Inches(5.55)
add_text(s_train, Inches(0.6), tbl_top, Inches(12.0), Inches(0.35),
         "Eval-set metrics  ·  populated by compare_training.py",
         size=14, bold=True, color=WHITE)

header = ["Regime", "AUC-ROC", "AUC-PR", "Rec@1%FPR", "Rec@10%FPR"]
rows = [
    ["A · SSL only",             "—",     "—",     "—",     "—"],
    ["B · Semi-supervised",      "TBD",   "TBD",   "TBD",   "TBD"],
    ["C · Fully supervised",     "TBD",   "TBD",   "TBD",   "TBD"],
]
col_widths = [Inches(4.05), Inches(2.05), Inches(2.05),
              Inches(2.05), Inches(2.05)]
tbl_y = tbl_top + Inches(0.40)
row_h = Inches(0.35)

# header row
x_cursor = Inches(0.6)
for w, h_label in zip(col_widths, header):
    cell = s_train.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, x_cursor, tbl_y, w, row_h)
    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
    cell.line.color.rgb = ACCENT; cell.line.width = Pt(0.5)
    add_text(s_train, x_cursor + Inches(0.10), tbl_y + Inches(0.05),
             w - Inches(0.20), row_h,
             h_label, size=11, bold=True, color=ACCENT)
    x_cursor += w

# data rows
for r_idx, row in enumerate(rows):
    y = tbl_y + row_h * (r_idx + 1)
    x_cursor = Inches(0.6)
    for c_idx, (w, val) in enumerate(zip(col_widths, row)):
        cell = s_train.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, x_cursor, y, w, row_h)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x10, 0x23, 0x40)
        cell.line.color.rgb = GREY; cell.line.width = Pt(0.25)
        add_text(s_train, x_cursor + Inches(0.10), y + Inches(0.05),
                 w - Inches(0.20), row_h,
                 val, size=11, bold=(c_idx == 0),
                 color=WHITE if c_idx == 0 else MUTED)
        x_cursor += w

add_footer(s_train,
           "Slide 3 / 4  •  Training regimes — see backend/.../compare_training.py")

# =====================================================================
# SLIDE 4 — Progress + Dates
# =====================================================================
s3 = prs.slides.add_slide(blank)
add_bg(s3)
add_accent_bar(s3, 0, 0, SW, Inches(0.12), color=ACCENT)

add_text(s3, Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.6),
         "Progress & Timeline", size=30, bold=True, color=WHITE)
add_text(s3, Inches(0.6), Inches(0.95), Inches(12.0), Inches(0.4),
         "Where we are — and what comes next.", size=15, color=ACCENT)

# ---- DONE column ----
col_top = Inches(1.65); col_h = Inches(4.3); col_w = Inches(4.05)
col_done_x = Inches(0.6)
done = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                           col_done_x, col_top, col_w, col_h)
done.fill.solid(); done.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
done.line.color.rgb = GREEN; done.line.width = Pt(1)

add_chip(s3, col_done_x + Inches(0.25), col_top + Inches(0.25),
         Inches(1.3), Inches(0.4), "DONE", fill=GREEN, color=NAVY, size=11)
add_bullets(s3, col_done_x + Inches(0.25), col_top + Inches(0.85),
            col_w - Inches(0.5), col_h - Inches(1.0),
            [
                "Ingestion layer (PCAP / Zeek / NetFlow) — canonical flow schema",
                "Temporal Graph Modeling — snapshots & edge features",
                "Project architecture & repository scaffolding",
                "Network discovery (baseline topology)",
            ],
            size=12, color=WHITE, bullet_color=GREEN, line_spacing=1.3)

# ---- IN PROGRESS column ----
col_ip_x = Inches(4.85)
ip = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                         col_ip_x, col_top, col_w, col_h)
ip.fill.solid(); ip.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
ip.line.color.rgb = AMBER; ip.line.width = Pt(1)

add_chip(s3, col_ip_x + Inches(0.25), col_top + Inches(0.25),
         Inches(1.8), Inches(0.4), "IN PROGRESS",
         fill=AMBER, color=NAVY, size=11)
add_bullets(s3, col_ip_x + Inches(0.25), col_top + Inches(0.85),
            col_w - Inches(0.5), col_h - Inches(1.0),
            [
                "Detection Core — TGNN model design & training loop",
                "Reconnaissance detector (heuristic baseline)",
                "Suspicion fusion module",
                "SOC frontend — Next.js dashboard scaffold",
            ],
            size=12, color=WHITE, bullet_color=AMBER, line_spacing=1.3)

# ---- TO DO column ----
col_td_x = Inches(9.1)
td = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                         col_td_x, col_top, col_w, col_h)
td.fill.solid(); td.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
td.line.color.rgb = ACCENT; td.line.width = Pt(1)

add_chip(s3, col_td_x + Inches(0.25), col_top + Inches(0.25),
         Inches(1.3), Inches(0.4), "TO DO", fill=ACCENT, color=NAVY, size=11)
add_bullets(s3, col_td_x + Inches(0.25), col_top + Inches(0.85),
            col_w - Inches(0.5), col_h - Inches(1.0),
            [
                "Correlation engine — recon→LM chaining",
                "Supervision layer — REST + WebSocket + alert engine",
                "Online pipeline orchestration (real-time)",
                "Full SOC frontend (live graph view)",
                "Evaluation on LANL + ablations",
            ],
            size=12, color=WHITE, bullet_color=ACCENT, line_spacing=1.3)

# ---- Timeline bar ----
tl_top = Inches(6.15)
add_text(s3, Inches(0.6), tl_top, Inches(6.0), Inches(0.35),
         "Key dates", size=14, bold=True, color=WHITE)

date_y = tl_top + Inches(0.45)

# date card 1
dc1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                          Inches(0.6), date_y, Inches(6.0), Inches(0.55))
dc1.fill.solid(); dc1.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
dc1.line.color.rgb = ACCENT; dc1.line.width = Pt(0.75)
add_text(s3, Inches(0.85), date_y + Inches(0.10), Inches(3.0), Inches(0.4),
         "Report submission", size=12, bold=True, color=MUTED)
add_text(s3, Inches(3.85), date_y + Inches(0.10), Inches(2.5), Inches(0.4),
         "TBD", size=14, bold=True, color=ACCENT, align=PP_ALIGN.RIGHT)

# date card 2
dc2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                          Inches(6.85), date_y, Inches(6.0), Inches(0.55))
dc2.fill.solid(); dc2.fill.fore_color.rgb = RGBColor(0x14, 0x2A, 0x4A)
dc2.line.color.rgb = GREEN; dc2.line.width = Pt(0.75)
add_text(s3, Inches(7.10), date_y + Inches(0.10), Inches(3.0), Inches(0.4),
         "Internship end date", size=12, bold=True, color=MUTED)
add_text(s3, Inches(10.10), date_y + Inches(0.10), Inches(2.5), Inches(0.4),
         "TBD", size=14, bold=True, color=GREEN, align=PP_ALIGN.RIGHT)

add_footer(s3, "Slide 4 / 4  •  Progress — dates to be confirmed before deposit")

# ---- Save ----
out = r"c:\Users\ilefb\OneDrive\Bureau\LateD\LateD_PFE2026_Ilef_Bennour.pptx"
prs.save(out)
print("WROTE:", out)
